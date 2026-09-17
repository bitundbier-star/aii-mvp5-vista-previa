"""Prueba manual de los cambios de esta ronda contra el servidor local."""
import requests

BASE = "http://127.0.0.1:8123"
s = requests.Session()

def check(nombre, cond):
    print(("OK  " if cond else "FAIL"), nombre)
    assert cond, nombre

# 1. Registro con monto tardío activado
r = s.post(f"{BASE}/registro", data={
    "nombre_conjunto": "Prueba Ronda 5",
    "direccion": "",
    "admin_nombre": "Ana Prueba",
    "login_email": "prueba_r5@test.com",
    "password": "clave123",
    "cuenta_email": "cuenta_r5@test.com",
    "monto_mensual": "800",
    "aplica_recargo_tardio": "1",
    "monto_mensual_tardio": "900",
    "monto_revision_meses": "12",
    "fecha_limite_pago": "11",
    "saldo_inicial": "0",
    "num_propiedades": "3",
}, allow_redirects=True)
check("registro con monto tardío -> 200 final", r.status_code == 200)
check("cayó en /propiedades", r.url.endswith("/propiedades?bienvenida=1"))

# 2. Descargar plantilla Excel
r = s.get(f"{BASE}/propiedades/plantilla.xlsx")
check("plantilla xlsx descarga", r.status_code == 200 and len(r.content) > 1000)
with open("/tmp/plantilla_descargada.xlsx", "wb") as f:
    f.write(r.content)

# 3. Editar propiedad 1 (numero cambia de 1 a 201) y ver que sigue en /propiedades sin error
r = s.post(f"{BASE}/propiedades/1/actualizar", data={
    "numero": "201", "tipo": "departamento", "nombre_dueno": "Juan Pérez",
    "celular_dueno": "", "email_dueno": "", "nombre_residente": "",
    "celular_residente": "", "email_residente": "", "notas": "",
    "tipo_saldo": "cero", "monto_saldo": "0",
}, allow_redirects=True)
check("actualizar propiedad numero->201", r.status_code == 200)
r = s.get(f"{BASE}/propiedades")
check("propiedades sigue ok tras renumerar", r.status_code == 200)
# La propiedad con numero 201 (id=1) debe seguir apareciendo ANTES que la 2 y 3
# en el HTML, porque ahora se ordena por id, no por numero.
pos201 = r.text.find('id="propiedad-1"')
pos2 = r.text.find('id="propiedad-2"')
check("orden estable: propiedad 1 (ahora #201) sigue antes que la 2", pos201 < pos2 and pos201 != -1)

# Completa las otras dos propiedades para poder pagar
for pid in (2, 3):
    s.post(f"{BASE}/propiedades/{pid}/actualizar", data={
        "numero": str(pid), "tipo": "casa", "nombre_dueno": f"Dueño {pid}",
        "celular_dueno": "", "email_dueno": "", "nombre_residente": "",
        "celular_residente": "", "email_residente": "", "notas": "",
        "tipo_saldo": "cero", "monto_saldo": "0",
    })

# 4. Monto sugerido antes de vencer (día actual probablemente antes del 11,
# dependiendo de la fecha real) - solo probamos que el endpoint responde.
r = s.get(f"{BASE}/pagos/monto-sugerido/1")
check("monto-sugerido responde JSON", r.status_code == 200 and "monto" in r.json())
print("   -> monto sugerido propiedad 1:", r.json())

# 5. Forzar fecha límite a un día ya pasado del mes actual, para que el mes
# actual cuente como vencido y probar el recargo tardío end-to-end.
import datetime as dt
hoy = dt.date.today()
dia_pasado = max(hoy.day - 1, 1) if hoy.day > 1 else 1
r = s.post(f"{BASE}/configuracion/fecha-limite", data={"fecha_limite_pago": str(dia_pasado)})
check("fecha limite actualizada", r.status_code in (200, 302))

r = s.get(f"{BASE}/pagos/monto-sugerido/1")
sugerido = r.json()["monto"]
print("   -> monto sugerido propiedad 1 (ya vencido, debería reflejar $900):", sugerido)
check("monto sugerido ahora incluye recargo tardio (900)", abs(sugerido - 900.0) < 0.01)

# 6. Registrar el pago con ese monto y revisar comprobante
r = s.post(f"{BASE}/pagos/nuevo", data={
    "propiedad_id": "1", "fecha_recepcion": hoy.isoformat(), "monto": str(sugerido),
    "concepto": "mantenimiento", "proyecto_id": "", "metodo_pago": "efectivo",
    "correo_destino": "",
}, allow_redirects=True)
check("pago registrado", r.status_code == 200)
check("comprobante muestra boton +Registrar pago", "+ Registrar pago" in r.text)
check("comprobante muestra Volver a pagos arriba y abajo", r.text.count("Volver a pagos") == 2)

# Imagen del comprobante: verificar que se genera sin error (bytes PNG)
import re
folio = re.search(r"AII-\d+-\d+", r.url).group(0)
img = s.get(f"{BASE}/comprobantes/{folio}/imagen")
check("imagen comprobante ok", img.status_code == 200 and img.content[:4] == b"\x89PNG")

# 7. Egresos: registrar uno, luego revisar que aparece en conceptos_previos
r = s.post(f"{BASE}/egresos/nuevo", data={"concepto": "Jardinería", "monto": "500", "fecha": hoy.isoformat()})
r = s.get(f"{BASE}/egresos")
check("egresos lista incluye datalist de conceptos previos", "Jardinería" in r.text and "conceptos-previos" in r.text)

# 8. Configuración: activar/editar monto tardío vía UI y recordatorios
r = s.post(f"{BASE}/configuracion/monto-tardio", data={
    "aplica_recargo_tardio": "1", "monto_mensual_tardio": "950",
})
check("guardar monto tardio en configuracion", r.status_code in (200, 302))
r = s.post(f"{BASE}/configuracion/recordatorios", data={"recordatorios_activos": "1"})
check("guardar recordatorios en configuracion", r.status_code in (200, 302))
r = s.get(f"{BASE}/configuracion")
check("configuracion refleja 950", "950.00" in r.text)

# 9. Importar Excel: modificar la plantilla descargada y subirla
# (se vuelve a descargar aquí porque la de antes es de cuando la propiedad
# todavía tenía el número 1, no el 201 que le pusimos después)
r = s.get(f"{BASE}/propiedades/plantilla.xlsx")
with open("/tmp/plantilla_descargada.xlsx", "wb") as f:
    f.write(r.content)
from openpyxl import load_workbook
wb = load_workbook("/tmp/plantilla_descargada.xlsx")
ws = wb["Propiedades"]
# La propiedad con numero "201" (antes 1) debería estar en el archivo.
fila_objetivo = None
for row in ws.iter_rows(min_row=2):
    if str(row[0].value).strip() == "201":
        fila_objetivo = row
        break
check("plantilla trae la propiedad 201", fila_objetivo is not None)
fila_objetivo[2].value = "Juan Pérez Actualizado desde Excel"
wb.save("/tmp/plantilla_modificada.xlsx")

with open("/tmp/plantilla_modificada.xlsx", "rb") as f:
    r = s.post(f"{BASE}/propiedades/importar", files={"archivo": ("plantilla.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
check("pantalla de confirmacion de importacion", r.status_code == 200 and "Así van a quedar" in r.text)
check("confirmacion muestra el nombre nuevo", "Juan Pérez Actualizado desde Excel" in r.text)

# Extraer el JSON del campo oculto y confirmar
import re as _re
m = _re.search(r'name="datos" value=\'(.*?)\'>', r.text, _re.S)
check("se encontro el campo datos", m is not None)
datos_json = m.group(1).replace("&#34;", '"').replace("&#39;", "'")
r = s.post(f"{BASE}/propiedades/importar/confirmar", data={"datos": datos_json}, allow_redirects=True)
check("confirmar importacion", r.status_code == 200)
check("mensaje de importadas", "actualizada" in r.text)
check("el nombre se aplico de verdad", "Juan Pérez Actualizado desde Excel" in r.text)

print("\nTODAS LAS PRUEBAS PASARON")
