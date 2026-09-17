"""Plantilla de Excel para el padrón de propiedades: descargarla, llenarla y
volver a subirla, en vez de capturar propiedad por propiedad en pantalla.

El archivo trae un renglón por cada propiedad que YA EXISTE en el conjunto
(las que se crearon numeradas al dar de alta la cuenta) — nunca renglones en
blanco. Al subirlo de regreso se empareja por «Número», que es el ancla
estable de cada propiedad, exactamente igual que en el resto de la app.
"""
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..models import TIPOS_PROPIEDAD, TIPOS_PROPIEDAD_DICT
from .cartera import orden_natural

SITUACIONES = ["No debe nada", "Debe esta cantidad", "Tiene saldo a favor"]
SITUACIONES_A_TIPO_SALDO = {
    "no debe nada": "cero",
    "debe esta cantidad": "debe",
    "tiene saldo a favor": "favor",
}

COLUMNAS = [
    ("Número", 12),
    ("Tipo de propiedad", 20),
    ("Nombre del propietario", 32),
    ("Celular del propietario", 22),
    ("Correo del propietario", 30),
    ("Nombre del residente", 32),
    ("Celular del residente", 22),
    ("Correo del residente", 30),
    ("Situación del saldo", 22),
    ("Monto del saldo", 18),
    ("Notas", 38),
]

AZUL = "1D4ED8"
GRIS_FONDO = "EFEFF2"
GRIS_BORDE = "D4D4D8"
FUENTE = "Arial"
_borde = Side(style="thin", color=GRIS_BORDE)
CAJA = Border(left=_borde, right=_borde, top=_borde, bottom=_borde)


def _situacion_de(propiedad) -> str:
    if propiedad.saldo_inicial > 0.005:
        return "Debe esta cantidad"
    if propiedad.saldo_inicial < -0.005:
        return "Tiene saldo a favor"
    return "No debe nada"


def generar_plantilla(conjunto) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Propiedades"
    ws.sheet_view.showGridLines = False

    for i, (enc, ancho) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=i, value=enc)
        c.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CAJA
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 30

    propiedades = sorted(conjunto.propiedades, key=orden_natural)
    for fila_idx, p in enumerate(propiedades, start=2):
        valores = [
            p.numero,
            p.tipo_legible,
            p.nombre_dueno or "",
            p.celular_dueno or "",
            p.email_dueno or "",
            p.nombre_residente or "",
            p.celular_residente or "",
            p.email_residente or "",
            _situacion_de(p),
            abs(round(p.saldo_inicial, 2)) if p.saldo_inicial else "",
            "" if (p.notas or "N/A") == "N/A" else p.notas,
        ]
        for i, valor in enumerate(valores, start=1):
            c = ws.cell(row=fila_idx, column=i, value=valor)
            c.font = Font(name=FUENTE, size=10)
            c.border = CAJA
            c.alignment = Alignment(vertical="center")
            if i == 10:
                c.number_format = '$#,##0.00'
        ws.cell(row=fila_idx, column=1).fill = PatternFill("solid", fgColor=GRIS_FONDO)

    ultima = len(propiedades) + 1

    dv_tipo = DataValidation(
        type="list", formula1='"%s"' % ",".join(e for _, e in TIPOS_PROPIEDAD), allow_blank=False
    )
    dv_tipo.error = "Elige un tipo de la lista."
    ws.add_data_validation(dv_tipo)
    dv_tipo.add(f"B2:B{ultima}")

    dv_sit = DataValidation(type="list", formula1='"%s"' % ",".join(SITUACIONES), allow_blank=False)
    dv_sit.error = "Elige una de las tres opciones de la lista."
    ws.add_data_validation(dv_sit)
    dv_sit.add(f"I2:I{ultima}")

    dv_monto = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv_monto.error = "El monto va siempre en positivo. Para saldo a favor, elige esa opción en «Situación del saldo»."
    ws.add_data_validation(dv_monto)
    dv_monto.add(f"J2:J{ultima}")

    ws.freeze_panes = "C2"
    if ultima >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{ultima}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def leer_plantilla(contenido: bytes, conjunto) -> tuple[list[dict], list[str]]:
    """Lee el archivo subido y devuelve (filas_validas, errores).

    Cada fila válida trae ya todo lo que `_guardar_ficha_propiedad` +
    `_saldo_desde_opciones` necesitan, más el id de la propiedad que le
    corresponde (encontrada por número). No escribe nada en la base de
    datos — eso lo hace la ruta que llama a esta función, después de que el
    administrador confirme en pantalla.
    """
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True)
    except Exception:
        return [], ["No se pudo abrir el archivo. ¿Seguro que es el .xlsx que descargaste de aquí?"]

    if "Propiedades" not in wb.sheetnames:
        return [], ["El archivo no tiene una hoja llamada «Propiedades». Usa la plantilla descargada de aquí, sin cambiarle el nombre a las hojas."]
    ws = wb["Propiedades"]

    por_numero = {(p.numero or "").strip().lower(): p for p in conjunto.propiedades}

    filas = []
    errores = []
    for fila in ws.iter_rows(min_row=2, values_only=False):
        valores = [c.value for c in fila]
        if all(v in (None, "") for v in valores):
            continue
        numero = str(valores[0] or "").strip()
        n_excel = fila[0].row
        if not numero:
            errores.append(f"Renglón {n_excel}: falta el número.")
            continue
        propiedad = por_numero.get(numero.lower())
        if not propiedad:
            errores.append(f"Renglón {n_excel}: no hay ninguna propiedad con el número «{numero}» en esta cuenta.")
            continue

        tipo_texto = str(valores[1] or "").strip()
        tipo_clave = next((k for k, v in TIPOS_PROPIEDAD_DICT.items() if v.lower() == tipo_texto.lower()), None)
        if not tipo_clave:
            errores.append(f"Renglón {n_excel} (número {numero}): «{tipo_texto}» no es un tipo de propiedad válido.")
            continue

        nombre_dueno = str(valores[2] or "").strip()
        if not nombre_dueno:
            errores.append(f"Renglón {n_excel} (número {numero}): falta el nombre del propietario.")
            continue

        situacion = str(valores[8] or "").strip().lower()
        tipo_saldo = SITUACIONES_A_TIPO_SALDO.get(situacion)
        if not tipo_saldo:
            errores.append(f"Renglón {n_excel} (número {numero}): «{valores[8]}» no es una situación de saldo válida.")
            continue
        try:
            monto_saldo = float(valores[9] or 0)
        except (TypeError, ValueError):
            errores.append(f"Renglón {n_excel} (número {numero}): el monto del saldo no es un número.")
            continue
        if monto_saldo < 0:
            errores.append(f"Renglón {n_excel} (número {numero}): el monto del saldo debe ir en positivo.")
            continue

        filas.append(
            {
                "propiedad_id": propiedad.id,
                "numero": propiedad.numero,
                "tipo": tipo_clave,
                "tipo_legible": TIPOS_PROPIEDAD_DICT[tipo_clave],
                "nombre_dueno": nombre_dueno,
                "celular_dueno": str(valores[3] or "").strip(),
                "email_dueno": str(valores[4] or "").strip(),
                "nombre_residente": str(valores[5] or "").strip(),
                "celular_residente": str(valores[6] or "").strip(),
                "email_residente": str(valores[7] or "").strip(),
                "tipo_saldo": tipo_saldo,
                "monto_saldo": round(monto_saldo, 2),
                "notas": str(valores[10] or "").strip(),
            }
        )

    return filas, errores
