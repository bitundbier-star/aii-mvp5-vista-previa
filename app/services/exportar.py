"""Descarga de todo el historial del conjunto, en un solo archivo ZIP.

Existe por una razón concreta: antes de borrar una cuenta hay que poder
llevarse lo que había dentro. El administrador en turno que aprieta "borrar" no
está borrando lo suyo — está borrando las cuentas de todos sus vecinos, incluidos
los comprobantes que ya les entregó. Ofrecerle la descarga convierte una
decisión abstracta en una concreta, y le deja algo en la mano si se arrepiente.

El ZIP trae:

* **historial.xlsx** — un solo Excel con una pestaña por tabla (propiedades,
  pagos, egresos, proyectos, cartera). Son los datos para seguir trabajando.
* **historial.html** — se abre de doble clic y se lee. Es la versión para
  imprimir, guardar o mandarle a la asamblea.
* **comprobantes/** — el PDF de cada comprobante de pago que se entregó.
* **archivos_egresos/** — los recibos, XML y comprobantes de cada egreso.
"""
import datetime as dt
import io
import os
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .cartera import estado_conjunto, orden_natural
from .formato import dinero, estado_saldo, mes_titulo
from .reportes import datos_reporte, meses_disponibles, saldo_acumulado


def _csv(encabezados, filas):
    """Antes cada tabla salía en un CSV; ahora cada una es una pestaña del
    mismo Excel. Se conserva el nombre para no tocar las funciones de abajo."""
    return (encabezados, filas)


def _num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return v


def _excel(hojas: list[tuple[str, tuple]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    encabezado = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="176BA0")
    for titulo, (encabezados, filas) in hojas:
        ws = wb.create_sheet(titulo[:31])
        ws.append(encabezados)
        for c in ws[1]:
            c.font = encabezado
            c.fill = fondo
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for fila in filas:
            ws.append(fila)
        for i, enc in enumerate(encabezados, start=1):
            largo = max([len(str(enc))] + [len(str(f[i - 1])) for f in filas[:300] if i - 1 < len(f)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 48)
            if any(k in enc.lower() for k in ("monto", "saldo", "recaudado", "esperado", "pagado", "fondo")):
                for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    celda.number_format = '$#,##0.00'
        ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _propiedades(conjunto):
    filas = []
    for p in sorted(conjunto.propiedades, key=orden_natural):
        est = estado_saldo(p.saldo_inicial)
        filas.append([
            p.numero, p.tipo_legible, p.nombre_dueno or "",
            p.celular_dueno or "", p.email_dueno or "",
            p.nombre_residente or "", p.celular_residente or "",
            p.email_residente or "", p.notas or "",
            _num(abs(p.saldo_inicial)), est["palabra"],
            "Sí" if p.activo else "No",
        ])
    return _csv([
        "Número", "Tipo", "Propietario", "Celular propietario", "Correo propietario",
        "Residente", "Celular residente", "Correo residente", "Notas",
        "Saldo inicial", "Situación del saldo inicial", "Activa",
    ], filas)


def _pagos(conjunto):
    filas = []
    for p in sorted(conjunto.pagos, key=lambda x: (x.fecha_recepcion, x.id)):
        filas.append([
            p.folio_recibo, p.folio, p.fecha_recepcion.isoformat(), p.propiedad.numero,
            p.propiedad.nombre_dueno or "", p.concepto_legible,
            p.proyecto.concepto if p.proyecto else "",
            p.metodo_pago_legible, _num(p.monto),
            "Sí" if (p.abona_a_cartera and not p.cancelado) else "No",
            "Sí" if p.cancelado else "No",
            p.cancelado_en.strftime("%Y-%m-%d %H:%M") if p.cancelado and p.cancelado_en else "",
            (p.concepto_descripcion or "Movimiento interno") if getattr(p, "interno", False) else "",
        ])
    return _csv([
        "Folio del comprobante", "Folio de la partida", "Fecha", "Propiedad", "Propietario",
        "Concepto", "Proyecto", "Método de pago", "Monto", "Baja la deuda", "Cancelado",
        "Cancelado el (UTC)", "Movimiento interno del sistema",
    ], filas)


def _egresos(conjunto):
    filas = [
        [
            e.fecha.isoformat(), e.concepto, e.forma_pago_legible, _num(e.monto),
            e.proyecto.concepto if e.proyecto else "",
            "Sí" if e.recibo_path else "No", "Sí" if e.xml_path else "No",
            "Sí" if e.comprobante_path else "No",
            "Sí" if getattr(e, "automatico", False) else "No",
        ]
        for e in sorted(conjunto.egresos, key=lambda x: (x.fecha, x.id))
    ]
    return _csv(
        ["Fecha", "Concepto", "Forma de pago", "Monto", "Proyecto", "Tiene recibo o factura", "Tiene XML",
         "Tiene comprobante de pago", "Registrado por el sistema"],
        filas,
    )


def _proyectos(conjunto):
    filas = []
    for p in sorted(conjunto.proyectos, key=lambda x: (x.fecha_alta, x.id)):
        if getattr(p, "eliminado", False):
            estado = "Eliminado"
        else:
            estado = p.estado_legible
        filas.append([
            p.concepto, p.descripcion or "", p.fecha_alta.isoformat(),
            p.fecha_limite_pago.isoformat() if p.fecha_limite_pago else "",
            estado, p.comentario_estado or "", p.financiamiento_legible,
            _num(p.monto_total), _num(p.monto_por_propiedad), _num(p.monto_del_fondo),
            _num(p.ajuste_centavos_fondo or 0), _num(p.total_recaudado),
            p.terminado_en.strftime("%Y-%m-%d") if p.terminado_en else "",
            p.cancelado_en.strftime("%Y-%m-%d") if p.cancelado and p.cancelado_en else "",
            (p.cancelado_motivo or "") if p.cancelado else "",
            p.eliminado_en.strftime("%Y-%m-%d") if getattr(p, "eliminado", False) and p.eliminado_en else "",
            (p.eliminado_motivo or "") if getattr(p, "eliminado", False) else "",
        ])
    return _csv([
        "Proyecto", "Descripción", "Fecha de alta", "Fecha límite", "Estado",
        "Comentario", "Financiamiento", "Monto total", "Monto por propiedad", "Monto del fondo",
        "Redondeo que absorbe el fondo", "Recaudado de vecinos",
        "Terminado el", "Cancelado el", "Motivo de cancelación", "Eliminado el", "Motivo de eliminación",
    ], filas)


def _cartera(conjunto):
    filas = []
    for e in estado_conjunto(conjunto):
        filas.append([
            e["propiedad"].numero, e["propiedad"].nombre_dueno or "",
            _num(e['total_esperado']), _num(e['total_pagado']),
            _num(abs(e['saldo'])), e["vista"]["palabra"],
        ])
    return _csv(
        ["Propiedad", "Propietario", "Esperado", "Pagado", "Saldo", "Situación"],
        filas,
    )


def _historial_html(conjunto) -> bytes:
    hoy = dt.date.today()
    estado = estado_conjunto(conjunto)
    meses = meses_disponibles(conjunto, todos=True)

    filas_cartera = "".join(
        f"<tr><td>{e['propiedad'].numero}</td>"
        f"<td>{e['propiedad'].nombre_dueno or ''}</td>"
        f"<td class='num'>{e['total_esperado']|0 if False else dinero(e['total_esperado'])}</td>"
        f"<td class='num'>{dinero(e['total_pagado'])}</td>"
        f"<td class='num {e['vista']['clase']}'>{e['vista']['monto']}</td>"
        f"<td>{e['vista']['palabra']}</td></tr>"
        for e in estado
    )

    filas_meses = ""
    for m in meses:
        r = datos_reporte(conjunto, m["anio"], m["mes"])
        filas_meses += (
            f"<tr><td>{r['periodo']}<br><small style='color:#666'>Administra {r.get('administrador', '')}</small></td>"
            f"<td class='num'>{dinero(r['saldo_apertura'])}</td>"
            f"<td class='num'>{dinero(r['ingresos_reales'])}</td>"
            f"<td class='num'>{dinero(r['total_egresos'])}</td>"
            f"<td class='num'><strong>{dinero(r['saldo_cierre'])}</strong></td></tr>"
        )

    html = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8">
<title>Historial de {conjunto.nombre}</title>
<style>
 body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
        color:#1a1a1a; max-width:900px; margin:0 auto; padding:36px 24px; line-height:1.6; }}
 h1 {{ font-size:26px; margin:0 0 4px; }}
 h2 {{ font-size:15px; text-transform:uppercase; letter-spacing:.05em; color:#666; margin:34px 0 10px; }}
 .sub {{ color:#666; margin:0 0 8px; }}
 table {{ width:100%; border-collapse:collapse; font-size:14px; }}
 th {{ text-align:left; color:#666; font-weight:600; border-bottom:1px solid #e5e5e5; padding:8px 6px; }}
 td {{ padding:7px 6px; border-bottom:1px solid #f0f0f0; }}
 .num {{ text-align:right; font-variant-numeric:tabular-nums; }}
 .saldo-debe {{ color:#b3261e; }} .saldo-favor, .saldo-cero {{ color:#0a7d3c; }}
 .nota {{ background:#EAF4FB; border-left:3px solid #176BA0; padding:12px 16px;
          border-radius:0 8px 8px 0; font-size:14px; }}
 footer {{ margin-top:36px; padding-top:16px; border-top:1px solid #e5e5e5; color:#999; font-size:12px; }}
</style></head><body>

<h1>{conjunto.nombre}</h1>
{f'<p class="sub">{conjunto.direccion}</p>' if conjunto.direccion else ''}
<p class="sub">Historial completo · Administrador en turno al descargar: {conjunto.admin_nombre}</p>
<p class="sub">Generado el {hoy.strftime('%d/%m/%Y')}</p>

<div class="nota">
  Esta es la copia del historial del conjunto al momento de descargarla.
  Junto a este archivo, en el mismo ZIP, vienen los datos en un Excel
  (historial.xlsx, una pestaña por tabla), el PDF de cada comprobante de pago
  y los recibos y facturas de los egresos.
</div>

<h2>Resumen</h2>
<table>
  <tr><td>Propiedades registradas</td><td class="num">{len(conjunto.propiedades)}</td></tr>
  <tr><td>Pagos registrados</td><td class="num">{len(conjunto.pagos)}</td></tr>
  <tr><td>Egresos registrados</td><td class="num">{len(conjunto.egresos)}</td></tr>
  <tr><td>Proyectos</td><td class="num">{len(conjunto.proyectos)}</td></tr>
  <tr><td>Cuenta abierta desde</td><td class="num">{conjunto.fecha_inicio_cobros.strftime('%d/%m/%Y')}</td></tr>
  <tr><td><strong>Saldo acumulado del conjunto</strong></td>
      <td class="num"><strong>{dinero(saldo_acumulado(conjunto, hoy))}</strong></td></tr>
</table>

<h2>Cartera por propiedad</h2>
<table>
  <tr><th>Propiedad</th><th>Propietario</th><th class="num">Esperado</th>
      <th class="num">Pagado</th><th class="num">Saldo</th><th>Situación</th></tr>
  {filas_cartera}
</table>

<h2>Mes a mes</h2>
<table>
  <tr><th>Mes</th><th class="num">Saldo de apertura</th><th class="num">Ingresos</th>
      <th class="num">Egresos</th><th class="num">Saldo al cierre</th></tr>
  {filas_meses or '<tr><td colspan="5">Todavía no hay meses cerrados.</td></tr>'}
</table>

<footer>
  Administración Inteligente de Inmuebles · Herramienta para vecinos que se auto-administran.<br>
  Los saldos se muestran siempre en positivo; la palabra de la última columna dice si es adeudo o saldo a favor.
</footer>
</body></html>"""
    return html.encode("utf-8")


def exportar_conjunto(conjunto) -> tuple[bytes, str]:
    """Arma el ZIP con todo el historial. Devuelve (contenido, nombre)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("historial.html", _historial_html(conjunto))
        z.writestr("historial.xlsx", _excel([
            ("Propiedades", _propiedades(conjunto)),
            ("Pagos", _pagos(conjunto)),
            ("Egresos", _egresos(conjunto)),
            ("Proyectos", _proyectos(conjunto)),
            ("Cartera", _cartera(conjunto)),
        ]))

        # El PDF de cada comprobante entregado (uno por recibo, con todas sus
        # partidas). Los movimientos internos del sistema no tienen comprobante.
        from .comprobante import comprobante_pdf

        recibos = {}
        for p in sorted(conjunto.pagos, key=lambda x: (x.fecha_recepcion, x.id)):
            if getattr(p, "interno", False):
                continue
            recibos.setdefault(p.folio_recibo, []).append(p)
        for folio, partidas in recibos.items():
            try:
                pdf = comprobante_pdf(partidas)
            except Exception:
                continue
            marca = " CANCELADO" if all(x.cancelado for x in partidas) else ""
            z.writestr(
                f"comprobantes/{partidas[0].fecha_recepcion.isoformat()} {folio} - {partidas[0].propiedad.etiqueta}{marca}.pdf",
                pdf,
            )

        # Los archivos de respaldo de cada egreso (recibo, XML, comprobante),
        # en una carpeta por egreso. Quien se lleva el historial se lleva
        # también las facturas.
        from ..database import DATA_DIR

        carpeta = os.path.join(DATA_DIR, "archivos_egresos")
        estaticos = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static")
        for e in sorted(conjunto.egresos, key=lambda x: (x.fecha, x.id)):
            for campo, etiqueta in (("recibo_path", "recibo"), ("xml_path", "xml"), ("comprobante_path", "comprobante de pago")):
                guardado = getattr(e, campo)
                if not guardado:
                    continue
                ruta = (
                    os.path.join(estaticos, guardado)
                    if guardado.startswith("egresos/")
                    else os.path.join(carpeta, os.path.basename(guardado))
                )
                if os.path.exists(ruta):
                    concepto = "".join(c for c in e.concepto if c.isalnum() or c in " -_").strip()[:40]
                    destino = f"archivos_egresos/{e.fecha.isoformat()} {concepto} ({e.id})/{etiqueta}{os.path.splitext(ruta)[1]}"
                    z.write(ruta, destino)

    limpio = "".join(
        c if c.isalnum() or c in " -_" else "" for c in conjunto.nombre
    ).strip().replace(" ", "_") or "conjunto"
    nombre = f"{limpio}_historial_{dt.date.today().isoformat()}.zip"
    return buffer.getvalue(), nombre
