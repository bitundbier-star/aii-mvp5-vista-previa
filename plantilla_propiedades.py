"""Plantilla de Excel para el padrón de propiedades: descargarla, llenarla y
volver a subirla, en vez de capturar propiedad por propiedad en pantalla.

El archivo trae un renglón por cada propiedad que YA EXISTE en el conjunto
(las que se crearon numeradas al dar de alta la cuenta) — nunca renglones en
blanco. Al subirlo de regreso se empareja por «ID», una columna interna que
no se puede editar y que identifica a la propiedad sin importar qué número
tenga puesto. El «Número» de al lado sí es editable: es la etiqueta que ve el
vecino (puede tener letras y números), y cambiarla aquí cambia el número real
de la propiedad, igual que si se editara a mano en Configuración.

Por qué separar las dos cosas: antes se emparejaba por «Número», así que si
alguien cambiaba el número de una propiedad y luego subía el Excel con ese
cambio, el sistema ya no encontraba a qué propiedad correspondía el renglón
(el número viejo ya no existía) y marcaba error. El ID nunca cambia, así que
el emparejamiento nunca se rompe, sin importar cuántas veces se edite el
número.
"""
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..models import TIPOS_PROPIEDAD, TIPOS_PROPIEDAD_DICT
from .cartera import orden_natural

SITUACIONES = ["No debe nada", "Debe esta cantidad", "Tiene saldo a favor"]
SITUACIONES_A_TIPO_SALDO = {
    "no debe nada": "cero",
    "debe esta cantidad": "debe",
    "tiene saldo a favor": "favor",
}

# El orden aquí es el orden real de las columnas en el Excel. ID va primero y
# es de solo lectura; todo lo demás es editable.
COLUMNAS = [
    ("ID", 8),
    ("Número", 12),
    ("Tipo de propiedad", 20),
    ("Nombre del propietario", 32),
    ("Celular del propietario", 22),
    ("Correo del propietario", 30),
    ("Nombre del residente", 32),
    ("Celular del residente", 22),
    ("Correo del residente", 30),
    ("Situación del saldo", 22),
    ("Monto del saldo", 18),
    ("Notas", 38),
]

# Posiciones (1-indexadas) de las columnas que si cambian de lugar hay que
# actualizar aquí y en `leer_plantilla`.
COL_ID = 1
COL_NUMERO = 2
COL_TIPO = 3
COL_NOMBRE_DUENO = 4
COL_CELULAR_DUENO = 5
COL_CORREO_DUENO = 6
COL_NOMBRE_RESIDENTE = 7
COL_CELULAR_RESIDENTE = 8
COL_CORREO_RESIDENTE = 9
COL_SITUACION = 10
COL_MONTO = 11
COL_NOTAS = 12

AZUL = "1D4ED8"
GRIS_FONDO = "EFEFF2"
GRIS_BORDE = "D4D4D8"
FUENTE = "Arial"
_borde = Side(style="thin", color=GRIS_BORDE)
CAJA = Border(left=_borde, right=_borde, top=_borde, bottom=_borde)

_LOCKED = Protection(locked=True)
_UNLOCKED = Protection(locked=False)


def _situacion_de(propiedad) -> str:
    if propiedad.saldo_inicial > 0.005:
        return "Debe esta cantidad"
    if propiedad.saldo_inicial < -0.005:
        return "Tiene saldo a favor"
    return "No debe nada"


def generar_plantilla(conjunto) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Propiedades"
    ws.sheet_view.showGridLines = False

    for i, (enc, ancho) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=i, value=enc)
        c.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CAJA
        c.protection = _LOCKED
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 30

    propiedades = sorted(conjunto.propiedades, key=orden_natural)
    for fila_idx, p in enumerate(propiedades, start=2):
        valores = [
            p.id,
            p.numero,
            p.tipo_legible,
            p.nombre_dueno or "",
            p.celular_dueno or "",
            p.email_dueno or "",
            p.nombre_residente or "",
            p.celular_residente or "",
            p.email_residente or "",
            _situacion_de(p),
            abs(round(p.saldo_inicial, 2)) if p.saldo_inicial else "",
            "" if (p.notas or "N/A") == "N/A" else p.notas,
        ]
        for i, valor in enumerate(valores, start=1):
            c = ws.cell(row=fila_idx, column=i, value=valor)
            c.font = Font(name=FUENTE, size=10)
            c.border = CAJA
            c.alignment = Alignment(vertical="center")
            # La columna ID se queda bloqueada (de solo lectura); todo lo
            # demás se puede editar aunque la hoja esté protegida.
            c.protection = _LOCKED if i == COL_ID else _UNLOCKED
            if i == COL_MONTO:
                c.number_format = '$#,##0.00'
        id_cell = ws.cell(row=fila_idx, column=COL_ID)
        id_cell.fill = PatternFill("solid", fgColor=GRIS_FONDO)

    ultima = len(propiedades) + 1

    dv_tipo = DataValidation(
        type="list", formula1='"%s"' % ",".join(e for _, e in TIPOS_PROPIEDAD), allow_blank=False
    )
    dv_tipo.error = "Elige un tipo de la lista."
    ws.add_data_validation(dv_tipo)
    dv_tipo.add(f"{get_column_letter(COL_TIPO)}2:{get_column_letter(COL_TIPO)}{ultima}")

    dv_sit = DataValidation(type="list", formula1='"%s"' % ",".join(SITUACIONES), allow_blank=False)
    dv_sit.error = "Elige una de las tres opciones de la lista."
    ws.add_data_validation(dv_sit)
    dv_sit.add(f"{get_column_letter(COL_SITUACION)}2:{get_column_letter(COL_SITUACION)}{ultima}")

    dv_monto = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv_monto.error = "El monto va siempre en positivo. Para saldo a favor, elige esa opción en «Situación del saldo»."
    ws.add_data_validation(dv_monto)
    dv_monto.add(f"{get_column_letter(COL_MONTO)}2:{get_column_letter(COL_MONTO)}{ultima}")

    ws.freeze_panes = f"{get_column_letter(COL_TIPO)}2"
    if ultima >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{ultima}"

    # Protege la hoja para que la columna ID no se pueda tocar por accidente,
    # sin bloquear la edición del resto de las columnas ni los filtros.
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def leer_plantilla(contenido: bytes, conjunto) -> tuple[list[dict], list[str]]:
    """Lee el archivo subido y devuelve (filas_validas, errores).

    Cada fila válida trae ya todo lo que `_guardar_ficha_propiedad` +
    `_saldo_desde_opciones` necesitan, más el id de la propiedad que le
    corresponde (encontrada por ID, no por número). No escribe nada en la
    base de datos — eso lo hace la ruta que llama a esta función, después de
    que el administrador confirme en pantalla.
    """
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True)
    except Exception:
        return [], ["No se pudo abrir el archivo. ¿Seguro que es el .xlsx que descargaste de aquí?"]

    if "Propiedades" not in wb.sheetnames:
        return [], ["El archivo no tiene una hoja llamada «Propiedades». Usa la plantilla descargada de aquí, sin cambiarle el nombre a las hojas."]
    ws = wb["Propiedades"]

    por_id = {p.id: p for p in conjunto.propiedades}

    candidatos = []
    errores = []
    for fila in ws.iter_rows(min_row=2, values_only=False):
        valores = [c.value for c in fila]
        if all(v in (None, "") for v in valores):
            continue
        n_excel = fila[0].row

        id_valor = valores[COL_ID - 1]
        try:
            id_int = int(id_valor)
        except (TypeError, ValueError):
            errores.append(
                f"Renglón {n_excel}: la columna «ID» no tiene un número válido. "
                "Esa columna no se debe modificar — solo sirve para identificar la propiedad."
            )
            continue
        propiedad = por_id.get(id_int)
        if not propiedad:
            errores.append(
                f"Renglón {n_excel}: no hay ninguna propiedad con ID «{id_int}» en esta cuenta. "
                "No modifiques la columna «ID»; si crees que esto es un error, vuelve a descargar la plantilla."
            )
            continue

        numero = str(valores[COL_NUMERO - 1] or "").strip()
        if not numero:
            errores.append(f"Renglón {n_excel} (ID {id_int}): falta el «Número» de la propiedad.")
            continue

        tipo_texto = str(valores[COL_TIPO - 1] or "").strip()
        tipo_clave = next((k for k, v in TIPOS_PROPIEDAD_DICT.items() if v.lower() == tipo_texto.lower()), None)
        if not tipo_clave:
            errores.append(f"Renglón {n_excel} (ID {id_int}): «{tipo_texto}» no es un tipo de propiedad válido.")
            continue

        nombre_dueno = str(valores[COL_NOMBRE_DUENO - 1] or "").strip()
        if not nombre_dueno:
            errores.append(f"Renglón {n_excel} (ID {id_int}): falta el nombre del propietario.")
            continue

        situacion = str(valores[COL_SITUACION - 1] or "").strip().lower()
        tipo_saldo = SITUACIONES_A_TIPO_SALDO.get(situacion)
        if not tipo_saldo:
            errores.append(f"Renglón {n_excel} (ID {id_int}): «{valores[COL_SITUACION - 1]}» no es una situación de saldo válida.")
            continue
        try:
            monto_saldo = float(valores[COL_MONTO - 1] or 0)
        except (TypeError, ValueError):
            errores.append(f"Renglón {n_excel} (ID {id_int}): el monto del saldo no es un número.")
            continue
        if monto_saldo < 0:
            errores.append(f"Renglón {n_excel} (ID {id_int}): el monto del saldo debe ir en positivo.")
            continue

        candidatos.append(
            {
                "n_excel": n_excel,
                "propiedad_id": propiedad.id,
                "numero": numero,
                "tipo": tipo_clave,
                "tipo_legible": TIPOS_PROPIEDAD_DICT[tipo_clave],
                "nombre_dueno": nombre_dueno,
                "celular_dueno": str(valores[COL_CELULAR_DUENO - 1] or "").strip(),
                "email_dueno": str(valores[COL_CORREO_DUENO - 1] or "").strip(),
                "nombre_residente": str(valores[COL_NOMBRE_RESIDENTE - 1] or "").strip(),
                "celular_residente": str(valores[COL_CELULAR_RESIDENTE - 1] or "").strip(),
                "email_residente": str(valores[COL_CORREO_RESIDENTE - 1] or "").strip(),
                "tipo_saldo": tipo_saldo,
                "monto_saldo": round(monto_saldo, 2),
                "notas": str(valores[COL_NOTAS - 1] or "").strip(),
            }
        )

    # Como el «Número» ahora se puede editar libremente, hay que asegurarse de
    # que dos propiedades no terminen con el mismo número — ni entre sí
    # dentro del archivo, ni con una propiedad que no vino en el archivo y se
    # queda con el número que ya tenía.
    ids_en_archivo = {c["propiedad_id"] for c in candidatos}
    por_numero_normalizado = {}
    for c in candidatos:
        por_numero_normalizado.setdefault(c["numero"].strip().lower(), []).append(c)

    ids_con_colision = set()
    for clave, filas_con_ese_numero in por_numero_normalizado.items():
        if len(filas_con_ese_numero) > 1:
            for c in filas_con_ese_numero:
                ids_con_colision.add(c["propiedad_id"])
                errores.append(
                    f"Renglón {c['n_excel']} (ID {c['propiedad_id']}): el número «{c['numero']}» "
                    "se repite en más de un renglón de este archivo."
                )
        for p in conjunto.propiedades:
            if p.id in ids_en_archivo:
                continue
            if (p.numero or "").strip().lower() == clave:
                for c in filas_con_ese_numero:
                    ids_con_colision.add(c["propiedad_id"])
                    errores.append(
                        f"Renglón {c['n_excel']} (ID {c['propiedad_id']}): el número «{c['numero']}» "
                        f"ya lo tiene otra propiedad de esta cuenta que no viene en el archivo (ID {p.id})."
                    )

    filas = [
        {k: v for k, v in c.items() if k != "n_excel"}
        for c in candidatos
        if c["propiedad_id"] not in ids_con_colision
    ]

    return filas, errores
